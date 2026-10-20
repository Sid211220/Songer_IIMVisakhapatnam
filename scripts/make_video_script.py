"""Write docs/VIDEO_SCRIPT.md from the video deck's speaker notes and docs/src/demo_narration.json.
    python scripts/make_video_script.py <path to Prahari_Video_Deck.pptx>"""
import json
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
deck = Presentation(sys.argv[1])
narr = json.loads((ROOT / "docs" / "src" / "demo_narration.json").read_text())
L = ["# Demo video script (about 20-21 minutes)", "",
     "Word-for-word narration for the submission video. It follows `Prahari_Video_Deck.pptx`: the presentation with the live dashboard demo embedded as slide 11. Speak at a relaxed pace, about 125-130 words a minute; a one-second pause between slides is fine.", "",
     "## How to record (PowerPoint, about 30 minutes of work)", "",
     "1. Open `Prahari_Video_Deck.pptx` in PowerPoint (Microsoft 365 or 2019 and later, Windows or Mac).",
     "2. Use earphones with a mic in a quiet room. Go to **Slide Show > Record** (or **Record > From Beginning**). The speaker notes show as a teleprompter; they are the same words as this script.",
     "3. Read each slide's notes, then press the right arrow.",
     "4. On slide 11 (*Live demo*), **click the video once to start it**, then read the demo narration below. Timestamps are video time, so glance at the video's timer and keep pace. If you are ahead, pause; if you fall behind, skip a sentence.",
     "5. Stop recording after the last slide. Then **File > Export > Create a Video > Full HD (1080p) > Use Recorded Timings and Narrations > Create Video**.",
     "6. Upload the MP4 to YouTube as **Unlisted**, title it *Prahari: Digital Twin Challenge 2026, Team Songer (IIM Visakhapatnam)*, and paste the link into `README.md` in place of `REPLACE_WITH_YOUR_UNLISTED_LINK`.", "",
     "If PowerPoint recording gives trouble, any screen recorder works: show the slides full screen, and for the demo part play `Prahari_Dashboard_Demo_silent.mp4` full screen while you read.", "", "---", ""]
t = 0.0
for i, s in enumerate(deck.slides, 1):
    notes = s.notes_slide.notes_text_frame.text
    texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    if any("Live demo" in x for x in texts):
        L += [f"## {int(t // 60)}:{int(t % 60):02d} · Slide {i} · Live demo (video, 7 minutes)", "", "*Click the video to start it. Times below are video time.*", ""]
        for ts, _, txt in narr:
            L += [f"**[{ts}]** {txt}", ""]
        t += 420
        continue
    head = "Title" if i == 1 else "Close" if i == len(deck.slides) else (texts[1] if len(texts) > 1 else texts[0])
    L += [f"## {int(t // 60)}:{int(t % 60):02d} · Slide {i} · {head}", "", notes, ""]
    t += len(notes.split()) / 128 * 60 + 3
L += ["---", "", f"Estimated total: about {int(t // 60)} minutes {int(t % 60)} seconds."]
(ROOT / "docs" / "VIDEO_SCRIPT.md").write_text("\n".join(L))
print(f"{t / 60:.1f} min")
